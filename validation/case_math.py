"""Conferência local determinística; este código não é executado pela arena."""

from __future__ import annotations

import argparse
import json
from dataclasses import dataclass
from datetime import datetime
from enum import Enum
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT / "starter-kit-arena-turim/material/retornos-mensais-2005-2026.xlsx"


class Currency(str, Enum):
    BRL = "BRL"
    USD = "USD"


class Unit(str, Enum):
    SINGLE = "unidade"
    MILLION = "milhoes"


class Period(str, Enum):
    STOCK = "saldo"
    MONTH = "mes"
    YEAR = "ano"


class Basis(str, Enum):
    NOMINAL = "nominal"
    REAL = "real"


@dataclass(frozen=True)
class Money:
    amount: float
    currency: Currency
    unit: Unit
    period: Period

    def __post_init__(self) -> None:
        if not isinstance(self.currency, Currency) or not isinstance(self.unit, Unit) or not isinstance(self.period, Period):
            raise TypeError("Money exige moeda, unidade e periodo tipados")

    @property
    def base_amount(self) -> float:
        return self.amount * (1_000_000 if self.unit == Unit.MILLION else 1)

    def as_millions(self) -> float:
        return self.base_amount / 1_000_000

    def __add__(self, other: Money) -> Money:
        if not isinstance(other, Money) or (self.currency, self.period) != (other.currency, other.period):
            raise TypeError("Soma exige moeda e periodo iguais")
        return Money(self.base_amount + other.base_amount, self.currency, Unit.SINGLE, self.period)

    def __sub__(self, other: Money) -> Money:
        if not isinstance(other, Money) or (self.currency, self.period) != (other.currency, other.period):
            raise TypeError("Subtracao exige moeda e periodo iguais")
        return Money(self.base_amount - other.base_amount, self.currency, Unit.SINGLE, self.period)


@dataclass(frozen=True)
class Rate:
    fraction: float
    currency: Currency
    period: Period
    basis: Basis
    after_tax: bool = False

    def __post_init__(self) -> None:
        if not isinstance(self.currency, Currency) or not isinstance(self.period, Period) or not isinstance(self.basis, Basis):
            raise TypeError("Rate exige moeda, periodo e base tipados")
        if self.fraction <= -1:
            raise ValueError("Taxa menor ou igual a -100%")


@dataclass(frozen=True)
class Month:
    date: datetime
    spx: float
    ibov: float
    usdbrl: float
    sofr: float
    cdi: float


def sale_proceeds(price: Money, cost: Money, tax_rate: float) -> tuple[Money, Money]:
    if price.period != Period.STOCK or cost.period != Period.STOCK:
        raise TypeError("Venda e custo devem ser saldos")
    gain = price - cost
    tax = Money(max(gain.base_amount, 0) * tax_rate, price.currency, Unit.SINGLE, Period.STOCK)
    return tax, price - tax


def real_after_tax_rate(gross: Rate, inflation: Rate, tax_rate: float) -> Rate:
    if gross.period != Period.YEAR or inflation.period != Period.YEAR:
        raise TypeError("Retorno e inflacao devem ser anuais")
    if gross.currency != inflation.currency or gross.basis != Basis.NOMINAL or inflation.basis != Basis.NOMINAL:
        raise TypeError("Retorno bruto e inflacao devem ser nominais e na mesma moeda")
    if gross.after_tax or inflation.after_tax:
        raise TypeError("Nao aplique imposto duas vezes")
    nominal_net = gross.fraction - max(gross.fraction, 0) * tax_rate
    real = (1 + nominal_net) / (1 + inflation.fraction) - 1
    return Rate(real, gross.currency, Period.YEAR, Basis.REAL, after_tax=True)


def perpetual_income(capital: Money, real_rate: Rate, fee: Money | None = None) -> Money:
    if capital.period != Period.STOCK or real_rate.period != Period.YEAR or real_rate.basis != Basis.REAL:
        raise TypeError("Renda perpetua exige saldo e taxa real anual")
    if capital.currency != real_rate.currency:
        raise TypeError("Capital e taxa precisam ter a mesma moeda")
    fee_amount = 0.0
    if fee is not None:
        if fee.currency != capital.currency or fee.period != Period.YEAR:
            raise TypeError("Custo da casa exige mesma moeda e periodo anual")
        fee_amount = fee.base_amount
    return Money(capital.base_amount * real_rate.fraction - fee_amount, capital.currency, Unit.SINGLE, Period.YEAR)


def read_months(path: Path) -> list[Month]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Retornos Mensais de Ativos"]
    header = [cell.value for cell in sheet[7]]
    expected = ["Data", "SPX – Retorno Total", "IBOV – Retorno Total", "USDBRL", "SOFR Proxy", "CDI"]
    if header != expected:
        raise ValueError(f"Colunas inesperadas na planilha: {header}")
    months: list[Month] = []
    for row in sheet.iter_rows(min_row=8, values_only=True):
        if not isinstance(row[0], datetime) or any(value is None for value in row[1:]):
            continue  # linha-base de dezembro de 2004
        months.append(Month(row[0], *(float(value) for value in row[1:])))
    if not months or any(left.date >= right.date for left, right in zip(months, months[1:])):
        raise ValueError("Serie mensal vazia ou fora de ordem")
    return months


def backtest(months: list[Month], currency: Currency, equity_weight: float, tax_rate: float) -> dict[str, float | str]:
    if not 0 <= equity_weight <= 1:
        raise ValueError("Peso da renda variavel deve estar entre 0 e 1")
    wealth = peak = year_start = 1.0
    max_drawdown = 0.0
    for index, month in enumerate(months):
        fixed, equity = (month.cdi, month.ibov) if currency == Currency.BRL else (month.sofr, month.spx)
        monthly_return = (1 - equity_weight) * fixed + equity_weight * equity
        wealth *= 1 + monthly_return  # pesos repostos mensalmente
        last_month_of_year = index + 1 < len(months) and months[index + 1].date.year != month.date.year
        if last_month_of_year:
            wealth -= max(wealth - year_start, 0) * tax_rate
            year_start = wealth
        peak = max(peak, wealth)
        max_drawdown = max(max_drawdown, 1 - wealth / peak)
    return {
        "currency": currency.value,
        "equity_weight_pct": 100 * equity_weight,
        "max_drawdown_pct": 100 * max_drawdown,
        "ending_wealth_for_1_initial": wealth,
        "first_month": months[0].date.strftime("%Y-%m"),
        "last_month": months[-1].date.strftime("%Y-%m"),
    }


def usd_return_in_brl(usd_return: float, usdbrl_return: float) -> float:
    return (1 + usd_return) * (1 + usdbrl_return) - 1


def case_report(path: Path, onshore_equity: float, offshore_equity: float, annual_fee_millions: float) -> dict:
    brl_stock = lambda amount: Money(amount, Currency.BRL, Unit.MILLION, Period.STOCK)
    price, acquisition, existing_liquidity = brl_stock(300), brl_stock(20), brl_stock(100)
    tax_rate = 0.15
    sale_tax, proceeds = sale_proceeds(price, acquisition, tax_rate)
    family_capital = existing_liquidity + proceeds
    daughter_capital = brl_stock(25)
    father_capital = family_capital - daughter_capital
    gross_onshore = Rate(0.12 * (1 - onshore_equity) + 0.13 * onshore_equity, Currency.BRL, Period.YEAR, Basis.NOMINAL)
    real_onshore = real_after_tax_rate(gross_onshore, Rate(0.04, Currency.BRL, Period.YEAR, Basis.NOMINAL), tax_rate)
    fee = Money(annual_fee_millions, Currency.BRL, Unit.MILLION, Period.YEAR)
    gross_offshore = Rate(0.04 * (1 - offshore_equity) + 0.09 * offshore_equity, Currency.USD, Period.YEAR, Basis.NOMINAL)
    real_offshore = real_after_tax_rate(gross_offshore, Rate(0.02, Currency.USD, Period.YEAR, Basis.NOMINAL), tax_rate)
    months = read_months(path)
    return {
        "source": str(path.relative_to(ROOT)) if path.is_relative_to(ROOT) else str(path),
        "case_facts_brl_millions": {
            "farm_value": price.as_millions(), "farm_acquisition_cost": acquisition.as_millions(),
            "farm_sale_tax": sale_tax.as_millions(), "net_sale_proceeds": proceeds.as_millions(),
            "family_liquidity_after_sale": family_capital.as_millions(),
            "father_after_internal_transfer": father_capital.as_millions(),
            "daughter_after_internal_transfer": daughter_capital.as_millions(),
            "illiquid_share_before_sale_pct": 100 * price.base_amount / (price + existing_liquidity).base_amount,
            "farm_income_to_family_spending_ratio": 10 / 10,
        },
        "illustrative_onshore_scenario": {
            "equity_weight_pct": 100 * onshore_equity,
            "nominal_gross_return_pct": 100 * gross_onshore.fraction,
            "real_after_tax_return_pct": 100 * real_onshore.fraction,
            "family_real_income_before_fee_brl_millions": perpetual_income(family_capital, real_onshore).as_millions(),
            "family_real_income_after_fee_brl_millions": perpetual_income(family_capital, real_onshore, fee).as_millions(),
            "annual_fee_assumed_brl_millions": annual_fee_millions,
        },
        "illustrative_offshore_return_usd": {
            "equity_weight_pct": 100 * offshore_equity,
            "nominal_gross_return_pct": 100 * gross_offshore.fraction,
            "real_after_tax_return_pct": 100 * real_offshore.fraction,
        },
        "backtest": [
            backtest(months, Currency.BRL, onshore_equity, tax_rate),
            backtest(months, Currency.USD, offshore_equity, tax_rate),
        ],
        "conventions": [
            "Pesos rebalanceados mensalmente; retornos da planilha em fracao mensal.",
            "Imposto de 15% sobre lucro positivo no fim de cada ano civil completo; sem compensacao de perdas.",
            "Meses de 2026 ainda nao foram tributados; drawdown sobre riqueza apos imposto quando cobrado.",
            "Drawdown onshore em BRL e offshore em USD; conversao cambial somente para visao consolidada em BRL.",
            "Renda perpetua ilustrativa: retorno real esperado constante, sem volatilidade, saques ou mudanca cambial.",
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--onshore-equity", type=float, default=0.0, help="peso entre 0 e 1")
    parser.add_argument("--offshore-equity", type=float, default=0.0, help="peso entre 0 e 1")
    parser.add_argument("--annual-fee-millions", type=float, default=0.0)
    args = parser.parse_args()
    if args.annual_fee_millions < 0:
        parser.error("custo anual não pode ser negativo")
    print(json.dumps(case_report(args.xlsx, args.onshore_equity, args.offshore_equity, args.annual_fee_millions), indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
